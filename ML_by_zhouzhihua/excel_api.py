#!/usr/bin/env python
# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.chart import (
BarChart,
AreaChart3D,
AreaChart,
Reference,
Series,
)
from copy import deepcopy

wb = Workbook()
ws = wb.active
ws.title = 'Summary'

for i in range(1,31):
    sheetName = 'June' + str(i)
    wb.create_sheet(sheetName) # 批量创建worksheet

rows = [
['Number', 'Batch 1', 'Batch 2'],
[2, 40, 30],
[3, 40, 25],
[4, 50, 30],
[5, 30, 10],
[6, 25, 5],
[7, 50, 10],
]
for row in rows:
    ws.append(row)
chart = AreaChart()
chart.title = "Area Chart"
chart.style = 13
chart.x_axis.title = 'Test'
chart.y_axis.title = 'Percentage'
cats = Reference(ws, min_col=1, min_row=1, max_row=7)
data = Reference(ws, min_col=2, min_row=1, max_col=3, max_row=7)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "A10")

rows = [
['Number', 'Batch 1', 'Batch 2'],
[2, 30, 40],
[3, 25, 40],
[4 ,30, 50],
[5 ,10, 30],
[6, 5, 25],
[7 ,10, 50],
]
ws1 = wb.get_sheet_by_name('June1')
for row in rows:
    ws1.append(row)

chart = AreaChart3D()
chart.title = "Area Chart"
chart.style = 13
chart.x_axis.title = 'Test'
chart.y_axis.title = 'Percentage'
chart.legend = None
cats = Reference(ws1, min_col=1, min_row=1, max_row=7)
data = Reference(ws1, min_col=2, min_row=1, max_col=3, max_row=7)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws1.add_chart(chart, "A10")

rows = [
('Number', 'Batch 1', 'Batch 2'),
(2, 10, 30),
(3, 40, 60),
(4, 50, 70),
(5, 20, 10),
(6, 10, 40),
(7, 50, 30),
]
ws2 = wb.get_sheet_by_name('June2')
for row in rows:
    ws2.append(row)

chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Bar Chart"
chart1.y_axis.title = 'Test number'
chart1.x_axis.title = 'Sample length (mm)'
data = Reference(ws2, min_col=2, min_row=1, max_row=7, max_col=3)
cats = Reference(ws2, min_col=1, min_row=2, max_row=7)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 4
ws2.add_chart(chart1, "A10")


chart2 = deepcopy(chart1)
chart2.style = 11
chart2.type = "bar"
chart2.title = "Horizontal Bar Chart"
ws2.add_chart(chart2, "G10")
chart3 = deepcopy(chart1)
chart3.type = "col"
chart3.style = 12
chart3.grouping = "stacked"
chart3.overlap = 100
chart3.title = 'Stacked Chart'
ws2.add_chart(chart3, "A27")
chart4 = deepcopy(chart1)
chart4.type = "bar"
chart4.style = 13
chart4.grouping = "percentStacked"
chart4.overlap = 100
chart4.title = 'Percent Stacked Chart'
ws2.add_chart(chart4, "G27")

wb.save('Add_sheet.xlsx')